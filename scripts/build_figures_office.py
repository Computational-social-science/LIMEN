#!/usr/bin/env python
"""build_figures_office.py -- DOCX and PDF editions of the Stage 2 figure set, captions in the document.

WHY NOT PANDOC
    pandoc is not on this host's PATH and the project's pandoc MCP is not connected in this session. Rather
    than skip two of the four formats the manuscript needs, this script uses tooling that IS present:
    `python-docx` for a real Word file, and a headless Edge/Chrome print for the PDF. The PDF therefore comes
    from the same rendering engine a reader's browser would use, not from a second typesetting chain that
    could disagree with the HTML.

CAPTIONS
    Read from `viz/figures/FIGURE_CAPTIONS.md` - the same file the HTML edition reads - so the four formats
    cannot drift apart. The captions are NOT burned into the images.
"""

from __future__ import annotations

import argparse
import pathlib
import re
import shutil
import subprocess
import sys

REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent
FIG_DIR = REPO_ROOT / "viz" / "figures"
CAPTIONS = FIG_DIR / "FIGURE_CAPTIONS.md"
OUT_DOCX = REPO_ROOT / "viz" / "figures_stage2.docx"
OUT_PDF = REPO_ROOT / "viz" / "figures_stage2.pdf"
HTML = REPO_ROOT / "viz" / "figures_stage2.html"

FIGURES = [
    ("fig1_stage2_hypotheses", "Figure 1. The three pre-registered hypotheses across the noise grid"),
    ("fig2_stage2_error_destination", "Figure 2. Where the errors go: the gate rejects most of what noise breaks"),
    ("fig3_stage2_gate_separation", "Figure 3. Why the gate separates better under noise"),
]

BROWSERS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
]


def parse_captions() -> dict[str, str]:
    text = CAPTIONS.read_text(encoding="utf-8")
    return {m.group(1): m.group(2).strip()
            for m in re.finditer(r"^## (\S+)\n\n(.*?)(?=\n## |\Z)", text, re.S | re.M)}


def plain(s: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", s)


def build_docx(caps: dict[str, str]) -> int:
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        r = subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "python-docx"],
                           capture_output=True, text=True)
        if r.returncode != 0:
            print(f"  [WARN] python-docx unavailable and could not be installed: {r.stderr.strip()[:120]}")
            return 0
        import docx
        from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(10.5)
    h = doc.add_heading("LIMEN Phase I — Stage 2 figures", level=1)
    p = doc.add_paragraph()
    r = p.add_run("Each figure is followed by its caption, set in the document's type as a manuscript sets it. "
                  "Captions are generated from the same data as the plots and are not part of the images.")
    r.italic = True
    r.font.size = Pt(9)

    for stem, title in FIGURES:
        png = FIG_DIR / f"{stem}.png"
        if not png.exists():
            print(f"  [WARN] missing {png}")
            continue
        doc.add_heading(title, level=2)
        doc.add_picture(str(png), width=Inches(6.4))
        cap = doc.add_paragraph()
        run = cap.add_run(plain(title).split(". ", 1)[1] + ". ")
        run.bold = True
        cap.add_run(plain(caps[stem]))
        cap.paragraph_format.space_after = Pt(18)
        for rr in cap.runs:
            rr.font.size = Pt(8.5)
    doc.save(OUT_DOCX)
    print(f"  wrote {OUT_DOCX.relative_to(REPO_ROOT)}  ({OUT_DOCX.stat().st_size // 1024} KB)")
    return 0


def build_pdf() -> int:
    if not HTML.exists():
        print(f"  [WARN] no HTML at {HTML}; PDF skipped")
        return 0
    exe = next((b for b in BROWSERS if pathlib.Path(b).exists()), None)
    if not exe:
        print("  [WARN] no Edge/Chrome found; PDF skipped")
        return 0
    prof = REPO_ROOT / "viz" / ".pdf_profile"
    cmd = [exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
           f"--user-data-dir={prof}", f"--print-to-pdf={OUT_PDF}",
           HTML.resolve().as_uri()]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    shutil.rmtree(prof, ignore_errors=True)
    if OUT_PDF.exists() and OUT_PDF.stat().st_size > 10_000:
        print(f"  wrote {OUT_PDF.relative_to(REPO_ROOT)}  ({OUT_PDF.stat().st_size // 1024} KB)")
        return 0
    print(f"  [WARN] PDF not produced: {(r.stderr or r.stdout or '').strip().splitlines()[:1]}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=(__doc__ or "").split("\n")[0])
    args = ap.parse_args()
    if not CAPTIONS.exists():
        print(f"  [FAIL] no captions at {CAPTIONS}")
        return 1
    caps = parse_captions()
    build_docx(caps)
    build_pdf()
    return 0


if __name__ == "__main__":
    sys.exit(main())